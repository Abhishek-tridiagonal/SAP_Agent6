"""Convert the consolidated Job Library into the SAP IA05 upload workbook.

Preserves rows 0-4 of the IA05 template verbatim (Mandatory / Optional / SAP
codes / field lengths / human column names) and clones per-cell styling from
the template's sample rows 6-8 onto every written data row.

Row expansion (Option B, nested):
  For each counter:
    1 HEADER row     TL Ind=A, Item Flag=H, GenTL=TaskListGroup, Grp Cntr=NN,
                     Header Description=counter header, Plant=cfg.plant.werks.
    1 INTN summary   Item Flag=I, Item No=0010, Control Key=INTN,
                     Operation Description=counter header, Duration=total.
    N SBOP sub-ops   Item Flag=I, Item No=0010, Sub-Operation=0001/0002/...,
                     Control Key=SBOP, Operation Description=op text,
                     Duration=per-step hours.

Work Center is mapped from each counter's dominant equipment Discipline
(from EquipmentClassMap sheet 1).

Public API:
    export(cfg) -> Path
"""

from copy import copy as _copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from config import Config, load_config
from utils import DEFAULT_WORKCENTER, DISCIPLINE_TO_WORKCENTER, load_job_library


# IA05 column index -> position in template header row (0-indexed).
COLS = {
    "TCODE":     0,
    "TL_IND":    1,
    "PLNNR":     2,   # GenTL
    "TPLNR":     3,   # Functional Location TL
    "EQUNR":     4,   # Equipment TL
    "PLNAL":     5,   # Grp Counter
    "KTEXT":     6,   # Header Description
    "VAGRP":     7,   # Planner Group
    "ITEM_FLAG": 8,
    "VORNR":     9,   # Item No
    "UVORN":    10,   # Sub-Operation
    "WERKS":    11,   # Plant
    "STEUS":    12,   # Control Key
    "ARBPL":    13,   # Work Center
    "LTXA1":    14,   # Operation Description
    "ANZZL":    15,   # No of Persons
    "DAUNO":    16,   # Duration in hours
    "DAUNE":    17,   # Unit of Duration
}


# =========================================================================
# Style capture / apply (preserve template colour coding)
# =========================================================================
def _capture_row_style(ws, row_idx):
    styles = {}
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row_idx, column=col)
        styles[col] = {
            "font":          _copy(cell.font),
            "fill":          _copy(cell.fill),
            "border":        _copy(cell.border),
            "alignment":     _copy(cell.alignment),
            "number_format": cell.number_format,
            "protection":    _copy(cell.protection),
        }
    return styles


def _apply_row_style(ws, row_idx, styles):
    for col, s in styles.items():
        cell = ws.cell(row=row_idx, column=col)
        cell.font          = _copy(s["font"])
        cell.fill          = _copy(s["fill"])
        cell.border        = _copy(s["border"])
        cell.alignment     = _copy(s["alignment"])
        cell.number_format = s["number_format"]
        cell.protection    = _copy(s["protection"])


# =========================================================================
# Flatten the loaded library into the counter list the writer walks
# =========================================================================
def _counters_from_library(data: dict) -> list[dict]:
    """load_job_library -> flat list of counter dicts, each with discipline
    already attached (utils.load_job_library computes dominant discipline)."""
    out = []
    for group, cs in data["counters"].items():
        for cnt in cs:
            out.append(cnt)
    return out


# =========================================================================
# Writer
# =========================================================================
def _write_workbook(template_path: Path, counters: list[dict], output: Path,
                    plant_werks: str) -> None:
    wb = load_workbook(template_path)
    ws = wb.active

    # Capture template sample-row styles BEFORE deleting them.
    style_H    = _capture_row_style(ws, 6) if ws.max_row >= 6 else None
    style_INTN = _capture_row_style(ws, 7) if ws.max_row >= 7 else None
    style_SBOP = _capture_row_style(ws, 8) if ws.max_row >= 8 else None

    data_start_row = 6
    if ws.max_row >= data_start_row:
        ws.delete_rows(data_start_row, ws.max_row - data_start_row + 1)

    row = data_start_row

    def put(col_key, value):
        if value is None or value == "":
            return
        ws.cell(row=row, column=COLS[col_key] + 1, value=value)

    def apply_style(style):
        if style:
            _apply_row_style(ws, row, style)

    for cnt in counters:
        group  = cnt["group"]
        ctr    = cnt["counter"]
        header = cnt["header"]
        ops    = cnt["operations"]
        wc     = DISCIPLINE_TO_WORKCENTER.get(cnt.get("discipline", ""), DEFAULT_WORKCENTER)
        total  = round(sum(o["duration_hrs"] for o in ops), 2)

        # 1) HEADER row  (A / H)
        put("TCODE", "IA05"); put("TL_IND", "A"); put("PLNNR", group)
        put("PLNAL", ctr); put("KTEXT", header[:40]); put("ITEM_FLAG", "H")
        put("WERKS", plant_werks); put("ARBPL", wc)
        apply_style(style_H); row += 1

        # 2) INTN summary row (I / VORNR 0010 / no UVORN)
        put("TCODE", "IA05"); put("PLNAL", ctr); put("ITEM_FLAG", "I")
        put("VORNR", "0010"); put("WERKS", plant_werks); put("STEUS", "INTN")
        put("ARBPL", wc); put("LTXA1", header[:40]); put("ANZZL", 1)
        put("DAUNO", total); put("DAUNE", "HR")
        apply_style(style_INTN); row += 1

        # 3) SBOP sub-op rows (I / VORNR 0010 / UVORN 0001, 0002, ...)
        for i, op in enumerate(ops, start=1):
            put("TCODE", "IA05"); put("PLNAL", ctr); put("ITEM_FLAG", "I")
            put("VORNR", "0010"); put("UVORN", f"{i:04d}")
            put("WERKS", plant_werks); put("STEUS", "SBOP"); put("ARBPL", wc)
            put("LTXA1", op["description"][:40])
            put("ANZZL", 1); put("DAUNO", op["duration_hrs"]); put("DAUNE", "HR")
            apply_style(style_SBOP); row += 1

    ws.column_dimensions[get_column_letter(COLS["KTEXT"] + 1)].width = 42
    ws.column_dimensions[get_column_letter(COLS["LTXA1"] + 1)].width = 42

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


# =========================================================================
# Public entry
# =========================================================================
def export(cfg: Config | None = None) -> Path:
    """Read cfg.job_library_consolidated and write the SAP IA05 workbook at
    cfg.sap_upload. Returns the output path."""
    cfg = cfg or load_config()
    print("[sap_exporter] export")
    print(f"[sap_exporter]   input (consolidated library): {cfg.job_library_consolidated}")
    print(f"[sap_exporter]   template (IA05 reference)   : {cfg.paths.ia05_template}")
    print(f"[sap_exporter]   output (SAP-ready workbook) : {cfg.sap_upload}")
    print(f"[sap_exporter]   plant WERKS                 : {cfg.plant.werks}\n")

    print(f"[sap_exporter] Step 1/3 : loading consolidated library ({cfg.job_library_consolidated.name})")
    data = load_job_library(cfg.job_library_consolidated)
    counters = _counters_from_library(data)
    print(f"[sap_exporter]            {len(counters)} counters loaded")

    print(f"[sap_exporter] Step 2/3 : preparing IA05 template "
          f"(preserving rows 0-4 and sample row styling H/INTN/SBOP)")
    print(f"[sap_exporter] Step 3/3 : writing counters to {cfg.sap_upload.name}")
    _write_workbook(cfg.paths.ia05_template, counters, cfg.sap_upload,
                    plant_werks=cfg.plant.werks)
    total_ops = sum(len(c["operations"]) for c in counters)
    total_rows = len(counters) * 2 + total_ops
    print(f"[sap_exporter]            {total_rows} data rows written "
          f"({len(counters)} H + {len(counters)} INTN + {total_ops} SBOP)")
    print(f"[sap_exporter] done -> {cfg.sap_upload}")
    return cfg.sap_upload


if __name__ == "__main__":
    export()
